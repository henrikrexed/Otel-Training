#!/usr/bin/env python3
"""Perform OTel training-structure working doc v2 (.docx) for Adam.

v2 rebuilds the structure ON TOP OF Adam's session plan (noise-generator,
four-stage collector pipeline, cold-open/replay) and reconciles it against
the original brief (declarative SDK config, BindPlane, the Perform-Hackathon
casino app). Advanced-focused. Collaboration draft, not a lab guide.
Issue: ISI-4025.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ACCENT = RGBColor(0x6A, 0x1B, 0x9A)
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xB0, 0x00, 0x20)
GREEN = RGBColor(0x1B, 0x5E, 0x20)

doc = Document()
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK
for lvl, sz, col in [("Heading 1", 16, ACCENT), ("Heading 2", 13, ACCENT), ("Heading 3", 11, DARK)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"; st.font.size = Pt(sz); st.font.color.rgb = col; st.font.bold = True


def para(text="", size=11, bold=False, italic=False, color=DARK, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size); r.bold = bold; r.italic = italic; r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def rich(parts, space_after=6):
    """parts: list of (text, dict-of-run-attrs)."""
    p = doc.add_paragraph()
    for text, attrs in parts:
        r = p.add_run(text)
        r.font.size = Pt(attrs.get("size", 11))
        r.bold = attrs.get("bold", False)
        r.italic = attrs.get("italic", False)
        r.font.color.rgb = attrs.get("color", DARK)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, level=0, bold_lead=None, color=DARK):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_lead:
        r = p.add_run(bold_lead); r.bold = True; r.font.color.rgb = color
    r2 = p.add_run(text); r2.font.color.rgb = color
    p.paragraph_format.space_after = Pt(3)
    return p


def table(rows, headers, widths=None, style="Light Grid Accent 4"):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = style
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]; c.text = ""
        run = c.paragraphs[0].add_run(h); run.bold = True; run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val)); run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def hr():
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(2); p.paragraph_format.space_after = Pt(8)
    pr = p._p.get_or_add_pPr(); pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6"); bottom.set(qn("w:space"), "1"); bottom.set(qn("w:color"), "6A1B9A")
    pbdr.append(bottom); pr.append(pbdr)


# ---------------- COVER ----------------
para("PERFORM · HOTDAY", size=11, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("OTel Beyond the Honeymoon", size=24, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Reconciled Training Structure — v2", size=14, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
para("Working draft for Adam & Henrik", size=12, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Built on Adam's session plan · reconciled against the original brief", size=10.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
para("3-hour instructor-led hands-on lab · one noise generator, one continuous collector pipeline", size=10.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
para("Owner: Henrik Rexed  ·  Issue: ISI-4025  ·  Status: co-development draft (v2)", size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("v1 was structured from the book/casino framing; v2 rebuilds on Adam's actual prep work.", size=9.5, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# ---------------- 0. WHAT CHANGED ----------------
doc.add_heading("0. What changed in v2, and how to read it", level=1)
para("You asked whether v1 considered Adam's prep work. It did not — v1 was structured from the OpenTelemetry "
     "Enterprise book and a casino-app framing, before Adam's session plan was on the table. v2 fixes that: it takes "
     "Adam's plan as the spine (you said you like the base of it) and reconciles it against the original brief.")
para("Adam's plan and the original brief genuinely disagree in three places. Those disagreements are real design "
     "forks, not wording — resolving them is the whole job of this document. Read Section 1 first; everything else "
     "flows from those three decisions.", italic=True, color=GREY)
hr()

# ---------------- 1. DECISIONS ----------------
doc.add_heading("1. Three forks to resolve before anything else", level=1)
para("Each row is a place where Adam's session plan and Henrik's original brief point in different directions. "
     "My recommendation is in the last column — but these are yours and Adam's to settle.", italic=True, color=GREY)

doc.add_heading("Fork A — Declarative SDK configuration: in or out?", level=2)
rich([("Brief: ", {"bold": True}),
      ("“cover the OTel SDK configured by files (declarative configuration)” — a named, core requirement.  ", {}),
      ("Adam: ", {"bold": True, "color": RED}),
      ("cut entirely — “spec too young, not mentioned at all now.”", {"color": RED})])
rich([("Recommendation: ", {"bold": True, "color": GREEN}),
      ("keep it, but tightly scoped — a 10–15 min framed demo (not a from-scratch lab), or as one of Adam's "
       "“increasingly difficult fast-finisher utilities.” It is the single most novel thing an advanced attendee "
       "takes home, and it is explicitly in the brief; dropping it silently would surprise Henrik. Scoping it small "
       "honours Adam's maturity concern.", {"color": GREEN})])

doc.add_heading("Fork B — BindPlane: hands-on control point, or a 5-minute aside?", level=2)
rich([("Brief: ", {"bold": True}),
      ("“utilize BindPlane for the collector configuration” — BindPlane is named as the vehicle for the collector work.  ", {}),
      ("Adam: ", {"bold": True, "color": RED}),
      ("BindPlane/OpAMP reduced to a 5-min scripted passage, no hands-on, because a separate BindPlane session "
       "exists and attendees are confirmed not to be in both.", {"color": RED})])
rich([("Recommendation: ", {"bold": True, "color": GREEN}),
      ("this is the #1 decision — it changes what attendees actually do for 90 minutes. Confirm first: is that "
       "separate BindPlane session at THIS event, and do attendees overlap? If BindPlane is meant to be THIS "
       "training's collector vehicle, run the four-stage pipeline THROUGH BindPlane instead of hand-written YAML — "
       "which makes Adam's OpAMP payoff land harder (“you just managed a fleet,” not “you wrote YAML by hand”). If "
       "the separate session truly owns BindPlane, Adam's 5-min aside is correct and we keep the block hand-authored.", {"color": GREEN})])

doc.add_heading("Fork C — Vehicle: the Perform-Hackathon casino app, or Adam's noise generator?", level=2)
rich([("Brief: ", {"bold": True}),
      ("“deliver the Perform-Hackathon-2026 application having issues in the configuration.”  ", {}),
      ("Adam: ", {"bold": True, "color": RED}),
      ("a purpose-built noise generator (minimal order-processing domain, two doors) — the highest-risk build item, "
       "built first.", {"color": RED})])
rich([("Recommendation: ", {"bold": True, "color": GREEN}),
      ("they are not really rivals — they serve different exercises. The noise generator is purpose-built for the "
       "collector block (redact / tame / enrich / sample) and for the cold-open→replay device; the hackathon app "
       "fits an SDK / declarative-config segment (broken SDK config to fix). If Fork A keeps declarative config, use "
       "the hackathon app for that short segment and the noise generator for the collector block. If build effort "
       "must be minimised, adopt the noise generator as the single vehicle (Adam's plan is more concrete) and drop "
       "the casino app.", {"color": GREEN})])

doc.add_heading("Two more that block a public title/date", level=2)
bullet("Title collision — Adam flags a submitted (REDACTED) talk with almost the same “honeymoon” framing. "
       "Resolve same-event? before the title is locked anywhere public: either differentiate or cross-promote.", color=RED)
bullet("Event & year — Adam's plan says “Perform 2027, Las Vegas”; the brief and repo say Perform-Hackathon-2026. "
       "Confirm which event this 3-hour Hotday is actually for.", color=RED)
hr()

# ---------------- 2. POSITIONING ----------------
doc.add_heading("2. Positioning (agreed direction)", level=1)
para("Adam and the brief agree on the altitude, so this is settled:")
for x in [
    "Audience: intermediate-to-advanced OTel users already running it in production — not first-timers. (Assumption to confirm: skill uniformity across the room.)",
    "No basic-OTel content: no “what is a span,” no first-instrumentation hello-world, no dashboard 101.",
    "The advanced core is the collector as a data-shaping control point: redaction, taming hostile sources, enrichment, and economically-responsible sampling — plus conventions enforcement (SemConv + Weaver + DT dictionary).",
    "Structural device: a cold open on raw, ugly telemetry; the same generator replayed clean at the end as the earned payoff.",
]:
    bullet(x)
para("Note on v1 vs v2: v1 over-indexed on declarative config + the book and under-indexed on collector data-shaping "
     "(redaction, hostile-source parsing, SemConv/Weaver). Adam's plan is stronger exactly there, so v2 leads with it.",
     italic=True, color=GREY)
hr()

# ---------------- 3. RUN OF SHOW ----------------
doc.add_heading("3. Run of show (Adam's plan, annotated)", level=1)
para("This is Adam's structure, preserved. Bracketed notes are where a Section-1 fork would change the row.", italic=True, color=GREY)
table(
    rows=[
        ["0:00–0:05", "Cold open — live dashboard on raw/ugly generator output. No slides.", "Demo"],
        ["0:05–0:20", "Framing — “does Dynatrace support X”; over- vs under-instrumentation, auto vs manual.", "Talk / Q&A"],
        ["0:20–1:50", "Hands-on block — one continuous pipeline, 4 stages (below).", "Lab — 90 min"],
        ["1:50–2:00", "Break (mandatory) + rescue buffer for anyone stuck on Stage 2 or 4.", "—"],
        ["2:00–2:05", "OpAMP / BindPlane in passing — scripted. [Fork B may make this hands-on]", "Talk"],
        ["2:05–2:20", "Semantic conventions + Weaver + DT dictionary — demo/discussion, enforce-at-scale framing.", "Demo"],
        ["2:20–2:30", "Replay the cold open — same dashboard, now clean. The earned payoff.", "Demo"],
        ["2:30–2:50", "Debrief — index cards, content-harvest: OTel limits (RUM, entities); telemetry for non-humans.", "Discussion"],
        ["2:50–3:00", "Wrap + buffer.", "—"],
    ],
    headers=["Time", "Segment", "Mode"],
    widths=[1.0, 4.7, 1.1],
)
para("")
doc.add_heading("The 90-minute hands-on block — four stages against one generator", level=2)
para("Config starting points: demo-opentelemetry-patterns scenarios 8, 14, 10, adapted to the generator's data "
     "(not the repo's sample data). [Fork B: done in BindPlane or hand-authored.]", italic=True, color=GREY)
for lead, rest in [
    ("Stage 1 — Redact.  ", "Strip the plaintext secret the generator plants in a log attribute."),
    ("Stage 2 — Tame the hostile door.  ", "Parse/normalise the legacy fulfillment feed (XML, rate-limited, HTTP-200-with-error-body, no TZ, etc.) into usable telemetry. The capstone — “you are the instrumentation now.”"),
    ("Stage 3 — Enrich.  ", "Join a service→owning-team lookup into the pipeline so the right humans get the right signals (analogue of demo-patterns scenario 10)."),
    ("Stage 4 — Sample.  ", "Build a tail-sampling policy: keep 100% errors + 100% slow, drop ~99% of healthy baseline. Facilitator narrates the live before/after volume drop — the “economically responsible sampling” payoff."),
]:
    bullet(rest, bold_lead=lead)
para("[Fork A — declarative SDK config, if kept — slots in here as a short framed demo or a fast-finisher utility, "
     "most naturally before Stage 1 or as an appendix track on the hackathon app.]", italic=True, color=ACCENT)
hr()

# ---------------- 4. TOPIC COVERAGE ----------------
doc.add_heading("4. Adam's topic list → where each lands", level=1)
table(
    rows=[
        ["“Does Dynatrace support X” (Technology + OTel)", "Framing (0:05–0:20)", "Talk / Q&A"],
        ["Over- vs under-instrumentation; auto vs manual (how/when/why)", "Framing", "Talk"],
        ["Collector: redaction, filtering, processing, sampling", "Hands-on block, Stages 1–4", "Lab — core"],
        ["Difficult systems: 3rd-party/legacy, non-friendly source data (XML/unstructured)", "Stage 2 (hostile door)", "Lab — capstone"],
        ["SemConv + Weaver + DT semantic dictionary", "2:05–2:20", "Demo / discussion"],
        ["Declarative configuration", "FORK A — see §1", "TBD"],
        ["OTel limitations: RUM, entities", "Debrief (2:30–2:50)", "Discussion"],
        ["Fleet management / OpAMP", "2:00–2:05 (or Fork B hands-on)", "Talk / TBD"],
    ],
    headers=["Adam's topic", "Where it lands", "Mode"],
    widths=[3.3, 2.3, 1.2],
)
hr()

# ---------------- 5. NOISE GENERATOR ----------------
doc.add_heading("5. The noise generator — highest-risk build item", level=1)
para("Adam's plan is explicit: build this first — the entire 90-minute block has nothing to work against until it "
     "exists. Summary of his spec (full spec carried verbatim in the appendix / repo):")
for lead, rest in [
    ("One source, two doors.  ", "Door A = native OTLP with production-shaped problems; Door B = the same events through a hostile legacy XML/HTTP feed. Both reflect the same underlying event stream."),
    ("Minimal, recognisable domain.  ", "A small order-processing service (Astroshop-shaped) — no new lore; attention goes to the problems, not the world."),
    ("Traffic that makes sampling real.  ", "~95% baseline, ~2–3% error bursts, ~2–3% slow bursts, so a keep-errors+slow / drop-baseline policy collapses volume by an order of magnitude — a dramatic live number."),
    ("Door A defects.  ", "Plaintext secret in a log attribute (Stage 1); missing ownership metadata + a service→team lookup (Stage 3); no sampling control (Stage 4)."),
    ("Door B nastiness.  ", "XML not JSON, rate-limited no-Retry-After, HTTP-200-with-error-body, local time no TZ, “3,214 units” strings, silent truncation, ALLCAPS_UNDERSCORE misspelled fields, occasional BOM."),
    ("Packaging.  ", "Single container, plain docker run, one instance per laptop; adjustable volume; seeded/deterministic mode; per-seat tag if output lands in a shared Dynatrace tenant."),
]:
    bullet(rest, bold_lead=lead)
para("Open build questions (Adam): impl language (Python vs Go); exact EPS/burst timing tuned in dry runs; Door B "
     "rate/row limits tight enough to force real handling without stalling the room; whether the container also drives "
     "the cold-open dashboard on startup.", italic=True, color=GREY)
hr()

# ---------------- 6. FAST FINISHERS & MECHANICS ----------------
doc.add_heading("6. Room mechanics (carried from Adam / the casino doc)", level=1)
for x in [
    "Fast-finisher utilities: 4–5 extra, increasingly-hard-to-tame little sources for attendees who fly through the config. Needed only if the room's skill level is mixed (open assumption). Declarative-config (Fork A) is a natural fast-finisher.",
    "Poker chips for hints — pending skill-level confirmation.",
    "Index cards for the debrief, repurposed as content harvesting (not open mic).",
    "Cold-open / replay as the structural payoff, both fed by the one generator.",
]:
    bullet(x)
para("Adam's read: most attendees will struggle with the basics of configuring a collector — so pace for the middle "
     "of the room, and use fast-finisher utilities as the pressure-release, not the main path.", italic=True, color=GREY)
hr()

# ---------------- 7. OPEN ITEMS ----------------
doc.add_heading("7. Open items to confirm before build", level=1)
for x in [
    "Section-1 forks A, B, C resolved (declarative config; BindPlane hands-on vs aside; vehicle).",
    "Title collision with the submitted (REDACTED) “honeymoon” talk — same event?",
    "Event & year (Perform 2026 vs 2027).",
    "Audience skill uniformity (decides the fast-finisher mechanic).",
    "Room + facilitator count for the 90-min hands-on block.",
    "demo-opentelemetry-patterns reuse terms for a live conference room.",
    "Script + rehearse the cold-open before/after so the replay lands clean.",
]:
    bullet(x)
hr()

# ---------------- 8. NEXT STEPS ----------------
doc.add_heading("8. Next steps", level=1)
for x in [
    "Adam + Henrik resolve the three forks in Section 1 (and the title/event opens).",
    "I fold the decisions into structure v3 and align docs/training-structure.md to the chosen shape.",
    "Only then: write the abstract — its wording depends entirely on Forks A/B/C.",
    "If the noise generator is adopted, it becomes the first build issue (highest risk) and blocks the hands-on labs.",
]:
    bullet(x)
para("")
para("Repo companions: docs/training-structure.md (to be re-aligned to this v2 shape), docs/abstract.md (draft — hold "
     "until forks resolve). This .docx is the shareable snapshot for Adam.", italic=True, color=GREY, size=9.5)

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "docs", "Perform-OTel-Training-Structure-for-Adam-v2.docx")
doc.save(out)
print("WROTE", out)
