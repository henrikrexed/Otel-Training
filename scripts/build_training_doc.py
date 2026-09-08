#!/usr/bin/env python3
"""Generate the Perform 2026 OTel training-structure working doc (.docx) for Adam.

Advanced-focused. This is a collaboration draft to co-develop the detailed
structure with Adam before writing the final abstract — NOT a lab guide.
Issue: ISI-4025.
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ACCENT = RGBColor(0x6A, 0x1B, 0x9A)   # deep purple
DARK = RGBColor(0x22, 0x22, 0x22)
GREY = RGBColor(0x66, 0x66, 0x66)

doc = Document()

# ---- base styles ----
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.font.color.rgb = DARK

for lvl, sz, col in [("Heading 1", 16, ACCENT), ("Heading 2", 13, ACCENT), ("Heading 3", 11, DARK)]:
    st = doc.styles[lvl]
    st.font.name = "Calibri"
    st.font.size = Pt(sz)
    st.font.color.rgb = col
    st.font.bold = True


def para(text="", size=11, bold=False, italic=False, color=DARK, align=None, space_after=6):
    p = doc.add_paragraph()
    if align:
        p.alignment = align
    r = p.add_run(text)
    r.font.size = Pt(size)
    r.bold = bold
    r.italic = italic
    r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p


def bullet(text, level=0, bold_lead=None):
    p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
        p.add_run(text)
    else:
        p.add_run(text)
    p.paragraph_format.space_after = Pt(3)
    return p


def kv_table(rows, headers, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 4"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ""
        run = hdr[i].paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(10)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(str(val))
            run.font.size = Pt(9.5)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t


def hr():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    pr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "6A1B9A")
    pbdr.append(bottom)
    pr.append(pbdr)


# =========================================================================
# COVER
# =========================================================================
para("PERFORM 2026 · HOTDAY", size=11, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("OpenTelemetry Beyond the Honeymoon", size=24, bold=True, color=DARK, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Detailed Training Structure — Advanced Track", size=14, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
para("Working draft for review with Adam", size=12, bold=True, color=ACCENT, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("3-hour instructor-led hands-on lab · BindPlane + Dynatrace · Vegas Casino app", size=10.5, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=18)
para("Owner: Henrik Rexed  ·  Issue: ISI-4025  ·  Status: co-development draft", size=10, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
para("Based on the book “Beyond the Honeymoon: OpenTelemetry at Enterprise Scale”", size=10, italic=True, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER)
doc.add_page_break()

# =========================================================================
# 0. HOW TO READ THIS
# =========================================================================
doc.add_heading("0. Purpose of this document", level=1)
para("This is a working draft, not the abstract and not a lab guide. Its job is to let Adam and Henrik "
     "agree on exactly what the 3 hours contain and at what depth, so the final abstract can be written "
     "with confidence. Everything here is up for negotiation.", italic=True, color=GREY)
para("Please read Sections 1 (Positioning) and 8 (Open decisions) first — that is where your input moves "
     "the design the most. The module breakdown (Section 4) is the current best proposal; treat the timings and "
     "the “advanced angle” of each module as the levers we tune together.", color=GREY)
hr()

# =========================================================================
# 1. POSITIONING — ADVANCED, NOT BASIC
# =========================================================================
doc.add_heading("1. Positioning: this is an advanced track", level=1)
para("The premise: the first OTel span is a honeymoon — one service, one collector, one dashboard, and it "
     "just works. This training starts the day after the honeymoon ends: polyglot services, a bill that "
     "arrives, and config sprawl. We do not teach “what is a span.” We teach how to run OTel as a fleet.")

doc.add_heading("What we explicitly do NOT cover (assumed knowledge)", level=2)
for x in [
    "What traces, metrics, and logs are, and why observability matters.",
    "Installing an SDK / writing a first manual span / “hello world” instrumentation.",
    "Basic dashboard reading — attendees have seen OTel data in a backend already.",
    "OTLP as a concept in the abstract — we assume they know it exists.",
]:
    bullet(x)
para("These are prerequisites we state up front so we can spend all 180 minutes on the hard parts.",
     italic=True, color=GREY)

doc.add_heading("What “advanced” means here (our differentiators)", level=2)
for lead, rest in [
    ("Declarative SDK configuration by file — ",
     "otel-config.yaml (file_format 1.1), not env-var soup. Most attendees have never seen this; it is the "
     "current frontier of SDK config and the reason to attend."),
    ("The collector as a control point, managed as a fleet — ",
     "processor ordering, routing, backpressure, and staged rollout via BindPlane / OpAMP — config as a "
     "versioned product, never SSH-edited."),
    ("Sampling as an explicit cost/fidelity strategy — ",
     "head vs. tail, tail-sampling policies, and keeping RED metrics alive under sampling with the spanmetrics connector."),
    ("Diagnosing real failure modes — ",
     "context-propagation breaks across languages, resource identity collisions, dropped signals — the "
     "things that actually page you, reproduced deliberately in the app."),
    ("Enterprise readiness discipline — ",
     "helm template → otelcol validate, Weaver semantic-convention checks, canary collector upgrades, persistent queues."),
]:
    bullet(rest, bold_lead=lead)
hr()

# =========================================================================
# 2. OUTCOMES
# =========================================================================
doc.add_heading("2. Learning outcomes", level=1)
para("By the end an attendee can:")
for x in [
    "Read and author OTel SDK declarative configuration (otel-config.yaml): resource detection, propagators, sampler trees, batch processors, OTLP exporters — and explain why env vars are ignored unless referenced with ${}.",
    "Diagnose the canonical enterprise failure modes from telemetry symptoms alone, and name the exact config surface that fixes each.",
    "Design a collector pipeline with correct processor order, enrichment, and routing — then validate and roll it out through BindPlane instead of editing YAML on a box.",
    "Choose and defend a sampling strategy as a cost/fidelity trade-off, and keep RED metrics accurate under sampling with spanmetrics.",
    "Apply the book’s enterprise readiness checklist to their own environment.",
]:
    bullet(x)
hr()

# =========================================================================
# 3. THE VEHICLE
# =========================================================================
doc.add_heading("3. The vehicle: a deliberately broken casino", level=1)
para("We ship the open-source Perform-Hackathon-2026 “Vegas Casino” app with its OTel configuration "
     "deliberately broken. Attendees diagnose, fix, tune, and operate its telemetry pipeline — SDK first, "
     "collector second — and watch each fix land in Dynatrace. The broken app is the spine of the whole session.")
para("Stack (all in a GitHub Codespace — no local setup):", bold=True)
for x in [
    "Polyglot services: Node.js, Python, Go, Java — plus Redis and Postgres.",
    "kind cluster, Dynatrace Operator + tenant, BindPlane instance, load/chaos generator.",
    "Declarative-config note: file-based SDK config is currently richest in the Java SDK (zero-code agent). "
    "The declarative-config lab therefore targets the Java Scoring Service; other services illustrate env-var "
    "vs. file config and the propagation/resource fixes.",
]:
    bullet(x)
para("Decision to confirm with Adam: is the Java-only scope for declarative config acceptable, or do we want a "
     "second language demoed even if less mature? (See Section 8.)", italic=True, color=GREY)
hr()

# =========================================================================
# 4. MODULE BREAKDOWN
# =========================================================================
doc.add_heading("4. Detailed module breakdown (180 min)", level=1)
kv_table(
    rows=[
        ["0:00–0:15", "M0 — The honeymoon problem + env check", "Talk"],
        ["0:15–0:45", "M1 — Diagnose the broken casino (failure modes)", "Lab"],
        ["0:45–1:15", "M2 — Declarative SDK configuration by file", "Lab"],
        ["1:15–1:25", "Break", "—"],
        ["1:25–2:05", "M3 — The collector as control point (BindPlane)", "Lab — centerpiece"],
        ["2:05–2:35", "M4 — Sampling as a cost/fidelity strategy", "Lab"],
        ["2:35–2:55", "M5 — Operating at scale + readiness checklist", "Talk + demo"],
        ["2:55–3:00", "Close & survey", "—"],
    ],
    headers=["Time", "Module", "Mode"],
    widths=[1.2, 4.3, 1.3],
)
para("")

modules = [
    ("M0 — The honeymoon problem (15 min, talk + env check)",
     "Frame the day: the lifecycle of an observability signal, and what breaks when one service becomes fifty.",
     "Advanced angle: skip the “what is OTel” intro entirely; open on the enterprise pain (cost, sprawl, "
     "broken traces) so the room knows this is not a 101.",
     ["5-min env check: Codespace up, kind healthy, casino deployed, Dynatrace connected, BindPlane login. Helpers circulate."]),
    ("M1 — Diagnose the broken casino (30 min, lab)",
     "Not a signals tour — a diagnosis clinic. Attendees generate load and hunt planted defects from symptoms.",
     "Advanced angle: assume the signals are known; focus on the failure taxonomy — propagation breaks, "
     "resource-identity collisions, dropped signals — and the config surface behind each.",
     ["Traces stop at the frontend → missing context propagation into a game service.",
      "Services show as unknown_service / duplicated → broken resource config.",
      "Logs missing entirely → logger provider not wired.",
      "Done when: attendee can name each defect class and the surface that fixes it."]),
    ("M2 — Declarative SDK configuration by file (30 min, lab)",
     "The headline advanced skill. otel-config.yaml, file_format 1.1, OTEL_CONFIG_FILE; env vars ignored unless "
     "referenced via ${}; schema = opentelemetry-configuration.",
     "Advanced angle: most attendees have never configured an SDK by file. This is the single most novel thing "
     "they take home.",
     ["Fix resource detection (service/host/process/container) + propagators (tracecontext + baggage).",
      "Fix the sampler tree: parent_based with a correct root policy.",
      "Fix batch processor + otlp_http exporter per-signal paths (/v1/traces|metrics|logs); try otlp_grpc.",
      "Done when: correct file-based config committed; Java Scoring Service telemetry visible and correctly named."]),
    ("M3 — The collector as control point, with BindPlane (40 min, lab — centerpiece)",
     "Pipeline anatomy (receivers → processors → exporters); canonical order enrichment → filtering "
     "→ aggregation → routing → export; config as a versioned, validated, staged product — fleet "
     "management with OpAMP.",
     "Advanced angle: the emotional core — stop hand-editing YAML on boxes. Start from blueprints, roll out "
     "through BindPlane, and understand why managed config wins.",
     ["Import/inspect the casino’s broken collector config in BindPlane.",
      "Fix processor order and the missing memory_limiter; add resource/enrichment processors.",
      "Add routing: game services vs. platform telemetry.",
      "Validate + staged rollout to kind via BindPlane; contrast with raw YAML (helm template → otelcol validate).",
      "Done when: healthy collector fleet in BindPlane; telemetry flows through managed config; attendee can defend managed vs. SSH."]),
    ("M4 — Sampling as a cost/fidelity strategy (30 min, lab)",
     "Head vs. tail; parent_based + traceidratio trade-offs; tail_sampling policies (keep errors + slow traces, "
     "drop health checks); spanmetrics keeps RED alive under sampling.",
     "Advanced angle: sampling is a strategy, not a failure — 100% is rarely the answer. Make the cost/fidelity "
     "trade-off explicit and measurable.",
     ["Run load/chaos; dial sampling down; observe in Dynatrace what survives.",
      "Add a tail-sampling policy + spanmetrics connector via BindPlane; re-measure.",
      "Done when: attendee can state their sampling policy and defend the trade-off."]),
    ("M5 — Operating at scale + readiness checklist (20 min, talk + demo)",
     "The book’s enterprise readiness checklist: config safety (helm template → otelcol validate + Weaver), "
     "canary upgrades, persistent queues, backpressure policy.",
     "Advanced angle: leave them with an operational discipline, not just configs — fleet ops, staged upgrades, "
     "and contributing blueprints back upstream.",
     ["Q&A + where to go next (the book, OTel blueprints)."]),
]
for title, concept, angle, points in modules:
    doc.add_heading(title, level=2)
    para(concept)
    para(angle, italic=True, color=ACCENT, space_after=4)
    for pt in points:
        bullet(pt)
    para("")
hr()

# =========================================================================
# 5. PLANTED-DEFECT INVENTORY
# =========================================================================
doc.add_heading("5. Planted-defect inventory (seed — to finalize against the repo)", level=1)
kv_table(
    rows=[
        ["1", "Game service (e.g. Roulette, Python)", "Propagation not configured → traces stop at frontend", "M1", "SDK propagators"],
        ["2", "Node services", "Missing/wrong service.name resource", "M1", "Resource attributes/detection"],
        ["3", "Node services", "Logger provider not wired → no logs", "M1", "logger_provider config"],
        ["4", "Java Scoring Service", "Broken otel-config.yaml (sampler/exporter)", "M2", "Declarative-config fixes"],
        ["5", "Collector", "Processor order wrong; no memory_limiter", "M3", "BindPlane pipeline fix"],
        ["6", "Collector", "No routing/enrichment; raw YAML sprawl", "M3", "BindPlane managed config"],
        ["7", "Collector", "No tail sampling / spanmetrics", "M4", "tail_sampling + spanmetrics"],
    ],
    headers=["#", "Where", "Defect", "Module", "Fix"],
    widths=[0.4, 1.9, 2.6, 0.7, 1.6],
)
para("Observability Agent is validating each of these against the actual repo; the final list may shift.",
     italic=True, color=GREY)
hr()

# =========================================================================
# 6. WHAT MAKES THIS DIFFERENT
# =========================================================================
doc.add_heading("6. Why this beats a standard OTel session", level=1)
for lead, rest in [
    ("A real broken app, not slides — ", "every concept is a fix the attendee makes and sees land in Dynatrace."),
    ("Declarative config — ", "the newest, least-taught SDK capability; the main reason an experienced attendee shows up."),
    ("BindPlane fleet management — ", "config-as-a-product, the way real teams operate collectors, not the YAML-on-a-box demo everyone has seen."),
    ("Cost is a first-class topic — ", "sampling framed as an economic decision, with spanmetrics protecting metric accuracy."),
    ("Book-backed — ", "the enterprise readiness checklist gives attendees something durable to take home."),
]:
    bullet(rest, bold_lead=lead)
hr()

# =========================================================================
# 7. LOGISTICS & RISKS
# =========================================================================
doc.add_heading("7. Logistics & risks", level=1)
for x in [
    "Staffing: 1 lead + 2 floating assistants for ~40 attendees; more TAs for larger rooms.",
    "Time contingencies: M4 compresses to ~15 min (demo-driven) if behind; M5 can drop the demo, never the checklist.",
    "Cold-start risk: Codespace pre-warm instruction email the day before; fallback pairing for anyone whose env fails.",
    "Prereqs on attendees: laptop + GitHub account; basic Kubernetes literacy assumed.",
]:
    bullet(x)
hr()

# =========================================================================
# 8. OPEN DECISIONS FOR ADAM
# =========================================================================
doc.add_heading("8. Open decisions — where we need Adam’s input", level=1)
para("These are the design questions whose answers shape the abstract. Henrik + Adam to resolve together.", italic=True, color=GREY)
kv_table(
    rows=[
        ["Depth vs. breadth", "Is 5 modules the right count, or do we cut one (e.g. fold M1 into M2) to go deeper on the collector/sampling centerpiece?"],
        ["Declarative-config scope", "Java-only for the file-config lab, or add a second language even if the SDK support is less mature?"],
        ["Level statement", "Do we call it “Advanced” outright, or “Intermediate–Advanced” to protect the room size? What prereqs do we hard-state?"],
        ["Sampling ambition", "Is tail-sampling + spanmetrics enough, or do we want a cost-comparison / $$ before-after moment to make the economics land?"],
        ["BindPlane centrality", "M3 is the centerpiece at 40 min — too much, too little? Should BindPlane also carry M4’s rollout?"],
        ["Book tie-in", "How hard do we lean on the book in the abstract — named throughout, or a single credibility line?"],
        ["Title", "“Beyond the Honeymoon: OpenTelemetry in Practice” vs. “From First Span to Production Pipeline” vs. an Adam idea."],
    ],
    headers=["Decision", "Question for us to settle"],
    widths=[1.9, 5.0],
)
hr()

# =========================================================================
# 9. NEXT STEPS
# =========================================================================
doc.add_heading("9. Next steps", level=1)
for x in [
    "Henrik + Adam review Sections 1 and 8; mark decisions.",
    "Fold decisions back into this structure (v2).",
    "Only then: write the final abstract from the agreed structure.",
    "In parallel (already in flight): Observability Agent validates the planted defects against the repo; Content Writer holds the lab guide until structure is locked.",
]:
    bullet(x)
para("")
para("Companion working files in the repo: docs/training-structure.md (source of truth), docs/abstract.md (draft abstract). "
     "This .docx is the shareable snapshot for Adam.", italic=True, color=GREY, size=9.5)

import os
out = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))),
                   "docs", "Perform2026-OTel-Training-Structure-for-Adam.docx")
doc.save(out)
print("WROTE", out)
